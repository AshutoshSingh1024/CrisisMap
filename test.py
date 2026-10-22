
from openpyxl import load_workbook
wb = load_workbook("Whitebox_Test_Cases_CrisisMap.xlsx")
ws = wb.active
for row in ws.iter_rows(min_row=2, max_col=10):
    for cell in row:
        if cell.value == "QA Engineer":
            cell.value = "Ashutosh"
wb.save("Whitebox_Test_Cases_CrisisMap.xlsx")